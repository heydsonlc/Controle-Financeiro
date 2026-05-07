from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from backend.services.ir_relatorio_service import IrRelatorioService
    from backend.services.patrimonio_empresarial_service import PatrimonioEmpresarialService
except ImportError:
    from services.ir_relatorio_service import IrRelatorioService
    from services.patrimonio_empresarial_service import PatrimonioEmpresarialService


class PatrimonioRelatorioService:
    ERRO_PERFIL = 'Relatorio patrimonial disponivel apenas no perfil Empresa.'
    MIME_EXCEL = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    @classmethod
    def gerar_excel(cls, filtros=None):
        perfil, bens = cls._bens_empresa(filtros)
        workbook = Workbook()
        ws_resumo = workbook.active
        ws_resumo.title = 'Resumo'
        cls._preencher_resumo(ws_resumo, perfil, bens, filtros or {})
        ws_bens = workbook.create_sheet('Bens')
        cls._preencher_bens(ws_bens, bens)

        arquivo = BytesIO()
        workbook.save(arquivo)
        arquivo.seek(0)
        return arquivo, cls._nome_arquivo(filtros or {}, 'xlsx')

    @classmethod
    def gerar_pdf(cls, filtros=None):
        perfil, bens = cls._bens_empresa(filtros)
        resumo = cls._resumo(bens)
        linhas = [
            'Relatorio de Bens Patrimoniais',
            f"Perfil financeiro: {perfil.nome if perfil else 'Empresa'}",
            f"Data de geracao: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
            f"Filtros aplicados: {cls._descrever_filtros(filtros or {})}",
            '',
            f"Total de bens: {len(bens)}",
            f"Valor total: {cls._moeda(resumo['valor_total'])}",
            f"Com lastro/documento: {resumo['com_lastro']}",
            f"Pendentes de documento: {resumo['pendentes']}",
            '',
            'Lista de bens:',
        ]
        for indice, bem in enumerate(bens, start=1):
            linhas.append(
                f"{indice}. {bem.get('nome') or '-'} | "
                f"{bem.get('categoria') or '-'} | "
                f"{cls._data(bem.get('data_aquisicao'))} | "
                f"{cls._moeda(bem.get('valor_aquisicao'))} | "
                f"{bem.get('status_documental_label') or '-'}"
            )
        if not bens:
            linhas.append('Nenhum bem encontrado para os filtros selecionados.')
        linhas.extend([
            '',
            'Relatorio gerencial. Validar classificacao patrimonial e documental antes de uso contabil.',
        ])
        return BytesIO(IrRelatorioService._pdf_simples(linhas)), cls._nome_arquivo(filtros or {}, 'pdf')

    @classmethod
    def _bens_empresa(cls, filtros):
        perfil_id, perfil, is_empresa = PatrimonioEmpresarialService.contexto()
        if not is_empresa:
            raise PermissionError(cls.ERRO_PERFIL)
        return perfil, PatrimonioEmpresarialService.listar_bens(filtros or {})

    @classmethod
    def _preencher_resumo(cls, ws, perfil, bens, filtros):
        resumo = cls._resumo(bens)
        linhas = [
            ('Perfil financeiro', perfil.nome if perfil else 'Empresa'),
            ('Data de geracao', datetime.now().strftime('%d/%m/%Y %H:%M')),
            ('Filtros aplicados', cls._descrever_filtros(filtros)),
            ('Total de bens', len(bens)),
            ('Valor total', resumo['valor_total']),
            ('Com lastro/documento', resumo['com_lastro']),
            ('Pendentes de documento', resumo['pendentes']),
            ('Sem documento', resumo['sem_documento']),
        ]
        ws.append(['Campo', 'Valor'])
        for linha in linhas:
            ws.append(list(linha))
        ws.append([])
        ws.append(['Categoria', 'Quantidade', 'Valor total'])
        for categoria, dados in resumo['por_categoria'].items():
            ws.append([categoria, dados['quantidade'], dados['valor']])
        cls._formatar_planilha(ws, money_columns={2, 3})

    @classmethod
    def _preencher_bens(cls, ws, bens):
        headers = [
            'Nome',
            'Codigo',
            'Categoria',
            'Data de aquisicao',
            'Valor de aquisicao',
            'Fornecedor',
            'Documento',
            'Status',
            'Status documental',
            'Centro de custo',
            'Localizacao',
            'Responsavel',
        ]
        ws.append(headers)
        for bem in bens:
            ws.append([
                bem.get('nome'),
                bem.get('codigo'),
                bem.get('categoria'),
                cls._data(bem.get('data_aquisicao')),
                float(cls._decimal(bem.get('valor_aquisicao'))),
                bem.get('fornecedor'),
                bem.get('documento_label'),
                bem.get('status_label'),
                bem.get('status_documental_label'),
                bem.get('centro_custo'),
                bem.get('localizacao'),
                bem.get('responsavel'),
            ])
        cls._formatar_planilha(ws, money_columns={5})

    @staticmethod
    def _formatar_planilha(ws, money_columns=None):
        money_columns = money_columns or set()
        header_fill = PatternFill('solid', fgColor='1D4ED8')
        header_font = Font(color='FFFFFF', bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            letra = column[0].column_letter
            largura = max(len(str(cell.value or '')) for cell in column) + 2
            ws.column_dimensions[letra].width = min(max(largura, 12), 42)
            if column[0].column in money_columns:
                for cell in column[1:]:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = 'R$ #,##0.00'

    @classmethod
    def _resumo(cls, bens):
        por_categoria = {}
        valor_total = Decimal('0')
        com_lastro = 0
        pendentes = 0
        sem_documento = 0
        for bem in bens:
            valor = cls._decimal(bem.get('valor_aquisicao'))
            valor_total += valor
            categoria = bem.get('categoria') or 'Sem categoria'
            por_categoria.setdefault(categoria, {'quantidade': 0, 'valor': Decimal('0')})
            por_categoria[categoria]['quantidade'] += 1
            por_categoria[categoria]['valor'] += valor
            status = str(bem.get('status_documental') or '').upper()
            if status in {'COM_LASTRO', 'COM_DOCUMENTO', 'VALIDADO'}:
                com_lastro += 1
            if status in {'SEM_DOCUMENTO', 'PENDENTE', 'AGUARDANDO_CONTADOR', 'DIVERGENTE'}:
                pendentes += 1
            if status == 'SEM_DOCUMENTO':
                sem_documento += 1
        return {
            'valor_total': valor_total,
            'com_lastro': com_lastro,
            'pendentes': pendentes,
            'sem_documento': sem_documento,
            'por_categoria': por_categoria,
        }

    @staticmethod
    def _descrever_filtros(filtros):
        itens = []
        for chave in ('categoria', 'status', 'documento', 'ano', 'busca'):
            valor = filtros.get(chave) if hasattr(filtros, 'get') else None
            if valor:
                itens.append(f'{chave}={valor}')
        return '; '.join(itens) if itens else 'Sem filtros adicionais'

    @staticmethod
    def _nome_arquivo(filtros, extensao):
        ano = filtros.get('ano') if hasattr(filtros, 'get') else None
        sufixo = str(ano or datetime.now().year)
        return f'Bens_Patrimoniais_Empresa_{sufixo}.{extensao}'

    @staticmethod
    def _data(valor):
        texto = str(valor or '').strip()
        if not texto:
            return None
        try:
            return datetime.strptime(texto[:10], '%Y-%m-%d').strftime('%d/%m/%Y')
        except ValueError:
            return texto[:10]

    @staticmethod
    def _decimal(valor):
        try:
            return Decimal(str(valor or 0))
        except Exception:
            return Decimal('0')

    @classmethod
    def _moeda(cls, valor):
        numero = cls._decimal(valor)
        return f"R$ {numero:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')
