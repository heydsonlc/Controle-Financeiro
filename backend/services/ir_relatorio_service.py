from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from backend.models import IrComprovante
    from backend.services.ir_documento_service import IrDocumentoService
    from backend.services.perfil_financeiro_service import PerfilFinanceiroService
except ImportError:
    from models import IrComprovante
    from services.ir_documento_service import IrDocumentoService
    from services.perfil_financeiro_service import PerfilFinanceiroService


class IrRelatorioService:
    MOEDA_FORMATO = 'R$ #,##0.00'

    @classmethod
    def gerar_excel(cls, filtros=None):
        contexto = cls._contexto()
        comprovantes = cls._coletar_comprovantes(filtros)
        wb = Workbook()
        wb.remove(wb.active)

        if contexto['empresa']:
            cls._aba_resumo_empresa(wb, contexto, comprovantes, filtros or {})
            cls._aba_documentos_empresa(wb, comprovantes)
            cls._aba_lastros(wb, comprovantes)
        else:
            cls._aba_resumo_pessoal(wb, contexto, comprovantes, filtros or {})
            cls._aba_comprovantes_pessoal(wb, comprovantes)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer, cls._nome_arquivo(contexto, filtros or {}, 'xlsx')

    @classmethod
    def gerar_pdf(cls, filtros=None):
        contexto = cls._contexto()
        comprovantes = cls._coletar_comprovantes(filtros)
        linhas = cls._linhas_pdf_empresa(contexto, comprovantes, filtros or {}) if contexto['empresa'] else cls._linhas_pdf_pessoal(contexto, comprovantes, filtros or {})
        return BytesIO(cls._pdf_simples(linhas)), cls._nome_arquivo(contexto, filtros or {}, 'pdf')

    @classmethod
    def resumo(cls, filtros=None):
        comprovantes = cls._coletar_comprovantes(filtros)
        return cls._resumo(comprovantes)

    @classmethod
    def _coletar_comprovantes(cls, filtros=None):
        filtros = filtros or {}
        cls._validar_filtros(filtros)
        comprovantes = IrDocumentoService.listar_comprovantes(filtros)

        dedutivel = str(filtros.get('dedutivel') or '').strip().lower()
        if dedutivel in {'true', '1', 'sim'}:
            comprovantes = [item for item in comprovantes if item.dedutivel is True]
        elif dedutivel in {'false', '0', 'nao'}:
            comprovantes = [item for item in comprovantes if item.dedutivel is False]

        if str(filtros.get('somente_validados') or '').strip().lower() in {'true', '1', 'sim'}:
            comprovantes = [item for item in comprovantes if item.status == 'VALIDADO']

        status_lastro = str(filtros.get('status_lastro') or '').strip().upper()
        if status_lastro and status_lastro != 'TODOS':
            comprovantes = [item for item in comprovantes if cls._vinculo_principal(item, 'status_lastro', 'SEM_DOCUMENTO') == status_lastro]

        natureza = str(filtros.get('natureza') or '').strip().upper()
        if natureza:
            comprovantes = [item for item in comprovantes if cls._vinculo_principal(item, 'natureza') == natureza]

        tipo_entidade = str(filtros.get('tipo_entidade') or '').strip().upper()
        if tipo_entidade:
            comprovantes = [item for item in comprovantes if cls._vinculo_principal(item, 'tipo_entidade') == tipo_entidade]

        if str(filtros.get('somente_pendentes') or '').strip().lower() in {'true', '1', 'sim'}:
            comprovantes = [
                item for item in comprovantes
                if cls._vinculo_principal(item, 'status_lastro', 'SEM_DOCUMENTO') in {'SEM_DOCUMENTO', 'PENDENTE', 'DIVERGENTE', 'AGUARDANDO_CONTADOR'}
                or item.status in {'PENDENTE_REVISAO', 'ERRO_LEITURA', 'IMPORTADO'}
            ]

        return comprovantes

    @staticmethod
    def _validar_filtros(filtros):
        ano = filtros.get('ano')
        if ano not in {None, ''}:
            try:
                int(ano)
            except (TypeError, ValueError) as exc:
                raise ValueError('Ano do relatorio invalido') from exc

    @staticmethod
    def _contexto():
        contexto = IrDocumentoService.obter_contexto()
        modo = contexto.get('modo') or 'IRPF'
        return {
            'perfil': contexto.get('perfil') or {},
            'modo': modo,
            'empresa': modo == 'DOCUMENTOS_FISCAIS_EMPRESA',
            'titulo': contexto.get('titulo') or 'Imposto de Renda',
            'gerado_em': datetime.now(),
        }

    @classmethod
    def _resumo(cls, comprovantes):
        valor_total = sum((Decimal(str(item.valor or 0)) for item in comprovantes), Decimal('0'))
        return {
            'total': len(comprovantes),
            'valor_total': valor_total,
            'validados': sum(1 for item in comprovantes if item.status == 'VALIDADO'),
            'pendentes': sum(1 for item in comprovantes if item.status in {'PENDENTE_REVISAO', 'IMPORTADO'}),
            'erros': sum(1 for item in comprovantes if item.status == 'ERRO_LEITURA'),
            'vinculados': sum(1 for item in comprovantes if item.vinculos.filter_by(ativo=True).first()),
            'aguardando_contador': sum(1 for item in comprovantes if cls._vinculo_principal(item, 'status_lastro') == 'AGUARDANDO_CONTADOR'),
        }

    @classmethod
    def _aba_resumo_pessoal(cls, wb, contexto, comprovantes, filtros):
        ws = wb.create_sheet('Resumo')
        resumo = cls._resumo(comprovantes)
        linhas = [
            ('Perfil financeiro', contexto['perfil'].get('nome')),
            ('Modo', 'Imposto de Renda'),
            ('Ano-calendario', filtros.get('ano') or 'Todos'),
            ('Total de comprovantes', resumo['total']),
            ('Valor total potencialmente dedutivel', float(cls._valor_dedutivel(comprovantes))),
            ('Quantidade validada', resumo['validados']),
            ('Quantidade pendente', resumo['pendentes']),
            ('Quantidade com erro de leitura', resumo['erros']),
            ('Data/hora de geracao', contexto['gerado_em'].strftime('%d/%m/%Y %H:%M')),
        ]
        for linha in linhas:
            ws.append(linha)
        ws.append([])
        ws.append(['Categoria IR', 'Quantidade', 'Valor total', 'Validados', 'Pendentes'])
        for row in cls._totais_por_categoria(comprovantes):
            ws.append(row)
        cls._formatar_planilha(ws)

    @classmethod
    def _aba_comprovantes_pessoal(cls, wb, comprovantes):
        ws = wb.create_sheet('Comprovantes')
        ws.append([
            'Ano-calendario', 'Data do documento', 'Prestador', 'CPF/CNPJ', 'Tomador',
            'Valor', 'Categoria IR', 'Categoria de Despesa', 'Dedutivel', 'Status',
            'Confianca', 'Origem da classificacao', 'Nome do arquivo', 'Observacoes',
        ])
        for item in comprovantes:
            ws.append([
                item.ano_calendario,
                cls._data(item.data_documento),
                item.prestador_nome,
                item.prestador_cpf_cnpj,
                item.tomador_nome,
                float(item.valor) if item.valor is not None else None,
                item.categoria_ir.nome if item.categoria_ir else None,
                item.categoria.nome if item.categoria else None,
                'Sim' if item.dedutivel else 'Nao' if item.dedutivel is False else '',
                item.status,
                item.confianca,
                item.origem_classificacao,
                item.arquivo.nome_arquivo if item.arquivo else None,
                item.observacoes,
            ])
        cls._formatar_planilha(ws)

    @classmethod
    def _aba_resumo_empresa(cls, wb, contexto, comprovantes, filtros):
        ws = wb.create_sheet('Resumo')
        resumo = cls._resumo(comprovantes)
        ws.append(('Perfil financeiro', contexto['perfil'].get('nome')))
        ws.append(('Modo', 'Documentos Fiscais e Lastro'))
        ws.append(('Ano/periodo', filtros.get('ano') or 'Todos'))
        ws.append(('Total de documentos', resumo['total']))
        ws.append(('Valor total documentado', float(resumo['valor_total'])))
        ws.append(('Documentos vinculados', resumo['vinculados']))
        ws.append(('Documentos sem vinculo', max(resumo['total'] - resumo['vinculados'], 0)))
        ws.append(('Aguardando contador', resumo['aguardando_contador']))
        ws.append(('Data/hora de geracao', contexto['gerado_em'].strftime('%d/%m/%Y %H:%M')))
        ws.append([])
        ws.append(['Categoria Fiscal', 'Natureza', 'Quantidade', 'Valor total', 'Com lastro', 'Sem lastro'])
        for row in cls._totais_por_categoria_natureza(comprovantes):
            ws.append(row)
        cls._formatar_planilha(ws)

    @classmethod
    def _aba_documentos_empresa(cls, wb, comprovantes):
        ws = wb.create_sheet('Documentos')
        ws.append([
            'Ano', 'Data', 'Prestador/Emissor', 'CPF/CNPJ', 'Valor', 'Categoria Fiscal',
            'Categoria de Despesa', 'Natureza', 'Status do documento', 'Status do lastro',
            'Nome do arquivo', 'Observacoes',
        ])
        for item in comprovantes:
            ws.append([
                item.ano_calendario,
                cls._data(item.data_documento),
                item.prestador_nome,
                item.prestador_cpf_cnpj,
                float(item.valor) if item.valor is not None else None,
                item.categoria_ir.nome if item.categoria_ir else None,
                item.categoria.nome if item.categoria else None,
                cls._vinculo_principal(item, 'natureza'),
                item.status,
                cls._vinculo_principal(item, 'status_lastro', 'SEM_DOCUMENTO'),
                item.arquivo.nome_arquivo if item.arquivo else None,
                item.observacoes,
            ])
        cls._formatar_planilha(ws)

    @classmethod
    def _aba_lastros(cls, wb, comprovantes):
        ws = wb.create_sheet('Lastros')
        ws.append([
            'Documento', 'Data', 'Valor', 'Tipo de entidade vinculada', 'ID/Descricao da entidade',
            'Tipo de vinculo', 'Natureza', 'Status do lastro', 'Observacoes',
        ])
        registros = 0
        for item in comprovantes:
            for vinculo in item.vinculos.filter_by(ativo=True).all():
                registros += 1
                ws.append([
                    item.arquivo.nome_arquivo if item.arquivo else f'Comprovante {item.id}',
                    cls._data(item.data_documento),
                    float(item.valor) if item.valor is not None else None,
                    vinculo.tipo_entidade,
                    vinculo.resumo_entidade or vinculo.entidade_id,
                    vinculo.tipo_vinculo,
                    vinculo.natureza,
                    vinculo.status_lastro,
                    vinculo.observacoes,
                ])
        if registros == 0:
            ws.append(['Nenhum vinculo de lastro cadastrado para os filtros selecionados.'])
        cls._formatar_planilha(ws)

    @classmethod
    def _totais_por_categoria(cls, comprovantes):
        grupos = defaultdict(lambda: {'qtd': 0, 'valor': Decimal('0'), 'validados': 0, 'pendentes': 0})
        for item in comprovantes:
            chave = item.categoria_ir.nome if item.categoria_ir else 'Sem categoria IR'
            grupos[chave]['qtd'] += 1
            grupos[chave]['valor'] += Decimal(str(item.valor or 0))
            grupos[chave]['validados'] += 1 if item.status == 'VALIDADO' else 0
            grupos[chave]['pendentes'] += 1 if item.status in {'PENDENTE_REVISAO', 'IMPORTADO'} else 0
        return [(nome, dados['qtd'], float(dados['valor']), dados['validados'], dados['pendentes']) for nome, dados in sorted(grupos.items())]

    @classmethod
    def _totais_por_categoria_natureza(cls, comprovantes):
        grupos = defaultdict(lambda: {'qtd': 0, 'valor': Decimal('0'), 'lastro': 0, 'sem': 0})
        for item in comprovantes:
            natureza = cls._vinculo_principal(item, 'natureza') or 'SEM_NATUREZA'
            chave = (item.categoria_ir.nome if item.categoria_ir else 'Sem categoria fiscal', natureza)
            grupos[chave]['qtd'] += 1
            grupos[chave]['valor'] += Decimal(str(item.valor or 0))
            if item.vinculos.filter_by(ativo=True).first():
                grupos[chave]['lastro'] += 1
            else:
                grupos[chave]['sem'] += 1
        return [(cat, nat, dados['qtd'], float(dados['valor']), dados['lastro'], dados['sem']) for (cat, nat), dados in sorted(grupos.items())]

    @classmethod
    def _linhas_pdf_pessoal(cls, contexto, comprovantes, filtros):
        resumo = cls._resumo(comprovantes)
        linhas = [
            'Relatorio de Comprovantes - Imposto de Renda',
            f"Perfil financeiro: {contexto['perfil'].get('nome') or ''}",
            f"Ano-calendario: {filtros.get('ano') or 'Todos'}",
            f"Gerado em: {contexto['gerado_em'].strftime('%d/%m/%Y %H:%M')}",
            '',
            f"Total de comprovantes: {resumo['total']}",
            f"Valor potencialmente dedutivel: {cls._moeda(cls._valor_dedutivel(comprovantes))}",
            f"Pendentes de revisao: {resumo['pendentes']}",
            f"Validados: {resumo['validados']}",
            '',
            'Totais por Categoria IR:',
        ]
        linhas.extend([f"- {cat}: {qtd} doc(s), {cls._moeda(valor)}" for cat, qtd, valor, *_ in cls._totais_por_categoria(comprovantes)])
        linhas.extend(['', 'Comprovantes:'])
        linhas.extend(cls._linhas_documentos_pdf(comprovantes, empresa=False))
        linhas.extend(['', 'Relatorio informativo. A dedutibilidade deve ser validada pelo contador/responsavel fiscal.'])
        return linhas

    @classmethod
    def _linhas_pdf_empresa(cls, contexto, comprovantes, filtros):
        resumo = cls._resumo(comprovantes)
        linhas = [
            'Relatorio de Documentos Fiscais e Lastro',
            f"Perfil financeiro: {contexto['perfil'].get('nome') or ''}",
            f"Ano/periodo: {filtros.get('ano') or 'Todos'}",
            f"Gerado em: {contexto['gerado_em'].strftime('%d/%m/%Y %H:%M')}",
            '',
            f"Documentos fiscais: {resumo['total']}",
            f"Valor documentado: {cls._moeda(resumo['valor_total'])}",
            f"Documentos com lastro: {resumo['vinculados']}",
            f"Documentos sem lastro: {max(resumo['total'] - resumo['vinculados'], 0)}",
            f"Aguardando contador: {resumo['aguardando_contador']}",
            '',
            'Totais por Categoria Fiscal/Natureza:',
        ]
        linhas.extend([f"- {cat} / {nat}: {qtd} doc(s), {cls._moeda(valor)}, lastro {lastro}, sem {sem}" for cat, nat, qtd, valor, lastro, sem in cls._totais_por_categoria_natureza(comprovantes)])
        linhas.extend(['', 'Documentos:'])
        linhas.extend(cls._linhas_documentos_pdf(comprovantes, empresa=True))
        pendentes = [item for item in comprovantes if cls._vinculo_principal(item, 'status_lastro', 'SEM_DOCUMENTO') in {'SEM_DOCUMENTO', 'DIVERGENTE', 'AGUARDANDO_CONTADOR', 'PENDENTE'}]
        linhas.extend(['', 'Pendencias:'])
        linhas.extend(cls._linhas_documentos_pdf(pendentes, empresa=True) or ['Nenhuma pendencia encontrada.'])
        linhas.extend(['', 'Relatorio gerencial. A classificacao fiscal/contabil deve ser validada pelo contador.'])
        return linhas

    @classmethod
    def _linhas_documentos_pdf(cls, comprovantes, empresa):
        linhas = []
        for item in comprovantes[:80]:
            partes = [
                cls._data(item.data_documento) or '-',
                item.prestador_nome or (item.arquivo.nome_arquivo if item.arquivo else 'Sem prestador'),
                item.categoria_ir.nome if item.categoria_ir else '-',
                cls._moeda(item.valor or 0),
                item.status,
            ]
            if empresa:
                partes.extend([cls._vinculo_principal(item, 'natureza') or '-', cls._vinculo_principal(item, 'status_lastro', 'SEM_DOCUMENTO')])
            linhas.append(' | '.join(str(parte) for parte in partes))
        return linhas

    @staticmethod
    def _formatar_planilha(ws):
        header_fill = PatternFill('solid', fgColor='1d4ed8')
        header_font = Font(color='FFFFFF', bold=True)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row == 1 or (cell.value in {'Categoria IR', 'Categoria Fiscal', 'Ano-calendario', 'Ano', 'Documento'}):
                    cell.fill = header_fill
                    cell.font = header_font
                if isinstance(cell.value, float) and cell.column_letter:
                    cell.number_format = IrRelatorioService.MOEDA_FORMATO
        ws.freeze_panes = 'A2'
        if ws.max_row > 1 and ws.max_column > 1:
            ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            letra = column[0].column_letter
            tamanho = max(len(str(cell.value or '')) for cell in column)
            ws.column_dimensions[letra].width = min(max(tamanho + 2, 12), 42)

    @staticmethod
    def _vinculo_principal(comprovante, campo, default=None):
        vinculo = comprovante.vinculos.filter_by(ativo=True).first()
        return getattr(vinculo, campo, default) if vinculo else default

    @staticmethod
    def _valor_dedutivel(comprovantes):
        return sum((Decimal(str(item.valor or 0)) for item in comprovantes if item.dedutivel is not False), Decimal('0'))

    @staticmethod
    def _data(valor):
        return valor.strftime('%d/%m/%Y') if valor else None

    @staticmethod
    def _moeda(valor):
        numero = Decimal(str(valor or 0))
        return f"R$ {numero:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

    @staticmethod
    def _nome_arquivo(contexto, filtros, extensao):
        ano = filtros.get('ano') or datetime.now().year
        perfil = re.sub(r'[^A-Za-z0-9_-]+', '_', contexto['perfil'].get('nome') or 'Perfil').strip('_')
        prefixo = 'Documentos_Fiscais' if contexto['empresa'] else 'IRPF'
        return f'{prefixo}_{ano}_{perfil}.{extensao}'

    @staticmethod
    def _pdf_simples(linhas):
        # Gerador PDF minimo para evitar dependencia nova no backend.
        safe_lines = []
        for linha in linhas:
            texto = str(linha or '').replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
            safe_lines.append(texto.encode('latin-1', 'replace').decode('latin-1')[:115])

        conteudo_linhas = ['BT', '/F1 10 Tf', '50 800 Td', '14 TL']
        for idx, linha in enumerate(safe_lines):
            if idx and idx % 52 == 0:
                conteudo_linhas.append('ET')
                conteudo_linhas.append('BT')
                conteudo_linhas.append('/F1 10 Tf')
                conteudo_linhas.append('50 800 Td')
                conteudo_linhas.append('14 TL')
            conteudo_linhas.append(f'({linha}) Tj')
            conteudo_linhas.append('T*')
        conteudo_linhas.append('ET')
        stream = '\n'.join(conteudo_linhas).encode('latin-1', 'replace')

        objetos = [
            b'1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n',
            b'2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n',
            b'3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >> endobj\n',
            b'4 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n',
            b'5 0 obj << /Length ' + str(len(stream)).encode('ascii') + b' >> stream\n' + stream + b'\nendstream endobj\n',
        ]
        pdf = BytesIO()
        pdf.write(b'%PDF-1.4\n')
        offsets = [0]
        for obj in objetos:
            offsets.append(pdf.tell())
            pdf.write(obj)
        xref = pdf.tell()
        pdf.write(f'xref\n0 {len(objetos) + 1}\n'.encode('ascii'))
        pdf.write(b'0000000000 65535 f \n')
        for offset in offsets[1:]:
            pdf.write(f'{offset:010d} 00000 n \n'.encode('ascii'))
        pdf.write(f'trailer << /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode('ascii'))
        return pdf.getvalue()
