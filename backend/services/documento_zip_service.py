from collections import Counter
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import mimetypes
import re
import unicodedata
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import or_

try:
    from backend.models import IrComprovante, IrComprovanteVinculo
    from backend.services.documento_empresarial_service import DocumentoEmpresarialService
    from backend.services.perfil_financeiro_service import PerfilFinanceiroService
except ImportError:
    from models import IrComprovante, IrComprovanteVinculo
    from services.documento_empresarial_service import DocumentoEmpresarialService
    from services.perfil_financeiro_service import PerfilFinanceiroService


class DocumentoZipService:
    ERRO_PERFIL = 'Pacote de documentos empresariais disponivel apenas no perfil Empresa.'
    ERRO_SEM_DOCUMENTOS = 'Nenhum documento encontrado para os filtros selecionados.'
    MOEDA_FORMATO = 'R$ #,##0.00'

    PASTAS_POR_TIPO = {
        'PAGAMENTO_REALIZADO': 'pagamentos_realizados',
        'NOTA_FISCAL_RECEBIDA': 'notas_fiscais_recebidas',
        'NOTA_FISCAL_EMITIDA': 'notas_fiscais_emitidas',
        'DOCUMENTO_OBRIGATORIO': 'documentos_obrigatorios',
        'DOCUMENTO_SOCIETARIO': 'societarios',
        'PROCURACAO_REPRESENTACAO': 'procuracoes_representacoes',
        'CONTRATO_INSTRUMENTO': 'contratos_instrumentos',
        'CERTIDAO_LICENCA': 'certidoes_licencas',
        'PATRIMONIO_IMOBILIZADO': 'patrimonio',
        'CONTABIL_FISCAL': 'contabil_fiscal',
        'OUTRO': 'outros',
    }

    TIPOS_CURTOS = {
        'PAGAMENTO_REALIZADO': 'PAGAMENTO',
        'NOTA_FISCAL_RECEBIDA': 'NF_RECEBIDA',
        'NOTA_FISCAL_EMITIDA': 'NF_EMITIDA',
        'DOCUMENTO_OBRIGATORIO': 'DOC_OBRIGATORIO',
        'DOCUMENTO_SOCIETARIO': 'SOC',
        'PROCURACAO_REPRESENTACAO': 'PROCURACAO',
        'CONTRATO_INSTRUMENTO': 'CONTRATO',
        'CERTIDAO_LICENCA': 'CERTIDAO',
        'PATRIMONIO_IMOBILIZADO': 'PATRIMONIO',
        'CONTABIL_FISCAL': 'CONTABIL',
        'OUTRO': 'DOC',
    }

    @classmethod
    def gerar_zip_documentos_empresa(cls, filtros=None, perfil_financeiro_id=None):
        perfil = cls._contexto_empresa(perfil_financeiro_id)
        filtros = cls._normalizar_filtros(filtros or {})
        documentos = cls.coletar_documentos_para_zip(filtros, perfil)
        if not documentos:
            raise ValueError(cls.ERRO_SEM_DOCUMENTOS)

        entradas = cls._montar_entradas(documentos)
        indice = cls.montar_indice_documentos(entradas, filtros, perfil)

        buffer = BytesIO()
        with ZipFile(buffer, 'w', ZIP_DEFLATED) as zip_file:
            zip_file.writestr('indice_documentos.xlsx', indice.getvalue())
            for entrada in entradas:
                cls.incluir_arquivo_no_zip(zip_file, entrada)

        buffer.seek(0)
        return buffer, cls._nome_pacote(filtros)

    @classmethod
    def coletar_documentos_para_zip(cls, filtros, perfil):
        query = IrComprovante.query.filter(IrComprovante.perfil_financeiro_id == perfil.id)

        ano = cls._parse_int(filtros.get('ano'))
        if ano:
            query = query.filter(IrComprovante.ano_calendario == ano)

        status = str(filtros.get('status') or '').strip().upper()
        if status and status not in {'TODOS', 'TODAS'}:
            query = query.filter(IrComprovante.status == status)

        categoria_ir_id = cls._parse_int(filtros.get('categoria_ir_id'))
        if categoria_ir_id:
            query = query.filter(IrComprovante.categoria_ir_id == categoria_ir_id)

        query = DocumentoEmpresarialService.aplicar_filtros_listagem(query, filtros)

        busca = str(filtros.get('busca') or '').strip()
        if busca:
            termo = f'%{busca}%'
            query = query.filter(
                or_(
                    IrComprovante.prestador_nome.ilike(termo),
                    IrComprovante.prestador_cpf_cnpj.ilike(termo),
                    IrComprovante.observacoes.ilike(termo),
                )
            )

        return query.order_by(IrComprovante.data_documento.desc(), IrComprovante.created_at.desc()).all()

    @classmethod
    def montar_indice_documentos(cls, entradas, filtros, perfil):
        wb = Workbook()
        wb.remove(wb.active)
        cls._aba_indice(wb, entradas)
        cls._aba_resumo(wb, entradas, filtros, perfil)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer

    @classmethod
    def montar_resumo_zip(cls, entradas, filtros):
        total_valor = sum((Decimal(str(item['valor'] or 0)) for item in entradas), Decimal('0'))
        por_tipo = Counter(item['tipo_label'] for item in entradas)
        por_status = Counter(item['status_documental_label'] for item in entradas)
        com_lastro = sum(1 for item in entradas if item['status_lastro'] in {'COM_DOCUMENTO', 'VALIDADO', 'COM_LASTRO'})
        return {
            'total_documentos': len(entradas),
            'total_valor': total_valor,
            'por_tipo': por_tipo,
            'por_status': por_status,
            'com_lastro': com_lastro,
            'sem_lastro': max(len(entradas) - com_lastro, 0),
            'filtros': filtros,
        }

    @classmethod
    def resolver_pasta_zip(cls, documento):
        metadata = DocumentoEmpresarialService.serializar_metadata(
            DocumentoEmpresarialService._metadata_ou_padrao(documento),
            documento,
        )
        tipo = str(metadata.get('tipo_documental') or 'OUTRO').upper()
        return cls.PASTAS_POR_TIPO.get(tipo, 'outros')

    @classmethod
    def gerar_nome_seguro(cls, documento, sequencial):
        metadata = DocumentoEmpresarialService.serializar_metadata(
            DocumentoEmpresarialService._metadata_ou_padrao(documento),
            documento,
        )
        tipo = str(metadata.get('tipo_documental') or 'OUTRO').upper()
        tipo_curto = cls.TIPOS_CURTOS.get(tipo, 'DOC')
        nome_base = (
            metadata.get('documento_label')
            or documento.prestador_nome
            or (documento.arquivo.nome_arquivo if documento.arquivo else None)
            or f'Documento {documento.id}'
        )
        nome_limpo = cls._sanitize_filename(nome_base)
        extensao = cls._extensao_arquivo(documento.arquivo)
        return f'{sequencial:03d}_{tipo_curto}_{nome_limpo}{extensao}'

    @classmethod
    def incluir_arquivo_no_zip(cls, zip_file, entrada):
        arquivo = entrada.get('arquivo')
        if not arquivo or not arquivo.conteudo:
            return
        zip_file.writestr(entrada['nome_zip'], arquivo.conteudo)

    @classmethod
    def _contexto_empresa(cls, perfil_financeiro_id=None):
        perfil_id = perfil_financeiro_id or PerfilFinanceiroService.obter_perfil_ativo_id()
        perfil = PerfilFinanceiroService.obter_perfil_por_id(perfil_id)
        if not perfil or str(perfil.tipo or '').upper() != 'EMPRESA':
            raise PermissionError(cls.ERRO_PERFIL)
        return perfil

    @classmethod
    def _normalizar_filtros(cls, filtros):
        dados = dict(filtros or {})
        if dados.get('categoria_fiscal') and not dados.get('categoria_ir_id'):
            dados['categoria_ir_id'] = dados.get('categoria_fiscal')
        status = str(dados.get('status') or '').strip().upper()
        if status and status not in {'TODOS', 'TODAS'} and not dados.get('status_documental'):
            codigos_status = {item['codigo'] for item in DocumentoEmpresarialService.STATUS_DOCUMENTAIS}
            if status in codigos_status:
                dados['status_documental'] = status
                dados.pop('status', None)
        return dados

    @classmethod
    def _montar_entradas(cls, documentos):
        entradas = []
        for indice, documento in enumerate(documentos, start=1):
            metadata = DocumentoEmpresarialService.serializar_metadata(
                DocumentoEmpresarialService._metadata_ou_padrao(documento),
                documento,
            )
            vinculo = documento.vinculos.filter_by(ativo=True).order_by(IrComprovanteVinculo.created_at.desc()).first()
            pasta = cls.resolver_pasta_zip(documento)
            arquivo = documento.arquivo
            nome_arquivo = cls.gerar_nome_seguro(documento, indice)
            nome_zip = f'{pasta}/{nome_arquivo}' if arquivo else 'Arquivo nao incluido'
            observacoes = cls._observacoes(documento, metadata, arquivo)
            entradas.append({
                'numero': indice,
                'documento': documento,
                'metadata': metadata,
                'arquivo': arquivo,
                'nome_zip': nome_zip,
                'nome_original': arquivo.nome_arquivo if arquivo else '',
                'documento_label': metadata.get('documento_label') or documento.prestador_nome or f'Documento #{documento.id}',
                'tipo': metadata.get('tipo_documental') or 'OUTRO',
                'tipo_label': metadata.get('tipo_documental_label') or 'Outro',
                'categoria_documental': metadata.get('categoria_documental_label') or 'Sem categoria',
                'categoria_fiscal': documento.categoria_ir.nome if documento.categoria_ir else '',
                'data_emissao': cls._data(metadata.get('data_emissao') or documento.data_documento),
                'data_validade': cls._data(metadata.get('data_validade')),
                'prestador_emissor': documento.prestador_nome or metadata.get('orgao_emissor') or '',
                'cpf_cnpj': documento.prestador_cpf_cnpj or '',
                'valor': Decimal(str(documento.valor or 0)) if documento.valor is not None else None,
                'status_documental': metadata.get('status_documental') or 'ATIVO',
                'status_documental_label': metadata.get('status_documental_label') or 'Ativo',
                'status_lastro': vinculo.status_lastro if vinculo else 'SEM_DOCUMENTO',
                'entidade_vinculada': cls._entidade_vinculada(vinculo),
                'observacoes': observacoes,
            })
        return entradas

    @classmethod
    def _aba_indice(cls, wb, entradas):
        ws = wb.create_sheet('Indice')
        headers = [
            'Nº',
            'Nome original',
            'Nome no ZIP',
            'Documento',
            'Tipo documental',
            'Categoria documental',
            'Categoria fiscal',
            'Data de emissao',
            'Data de validade',
            'Prestador / emissor',
            'CPF/CNPJ',
            'Valor',
            'Status documental',
            'Status fiscal/lastro',
            'Entidade vinculada',
            'Observacoes',
        ]
        ws.append(headers)
        for entrada in entradas:
            ws.append([
                entrada['numero'],
                entrada['nome_original'],
                entrada['nome_zip'],
                entrada['documento_label'],
                entrada['tipo_label'],
                entrada['categoria_documental'],
                entrada['categoria_fiscal'],
                entrada['data_emissao'],
                entrada['data_validade'],
                entrada['prestador_emissor'],
                entrada['cpf_cnpj'],
                float(entrada['valor']) if entrada['valor'] is not None else None,
                entrada['status_documental_label'],
                entrada['status_lastro'],
                entrada['entidade_vinculada'],
                entrada['observacoes'],
            ])
        cls._formatar_planilha(ws)
        for cell in ws['L'][1:]:
            cell.number_format = cls.MOEDA_FORMATO

    @classmethod
    def _aba_resumo(cls, wb, entradas, filtros, perfil):
        ws = wb.create_sheet('Resumo')
        resumo = cls.montar_resumo_zip(entradas, filtros)
        linhas = [
            ('Perfil financeiro', perfil.nome),
            ('Data/hora de geracao', datetime.now().strftime('%d/%m/%Y %H:%M')),
            ('Ano-calendario', filtros.get('ano') or 'Todos'),
            ('Filtros aplicados', cls._descricao_filtros(filtros)),
            ('Total de documentos', resumo['total_documentos']),
            ('Total de valor documentado', float(resumo['total_valor'])),
            ('Total com lastro', resumo['com_lastro']),
            ('Total sem lastro', resumo['sem_lastro']),
            ('Aviso', 'Pacote gerencial para apoio ao contador. Conferir com a contabilidade antes de entregar obrigacoes fiscais.'),
        ]
        for linha in linhas:
            ws.append(linha)

        ws.append([])
        cls._adicionar_contagem(ws, 'Total por tipo documental', resumo['por_tipo'])
        cls._adicionar_contagem(ws, 'Total por status', resumo['por_status'])
        cls._formatar_planilha(ws)
        ws['B6'].number_format = cls.MOEDA_FORMATO

    @staticmethod
    def _adicionar_contagem(ws, titulo, dados):
        ws.append([titulo])
        ws.append(['Nome', 'Quantidade'])
        if dados:
            for nome, quantidade in sorted(dados.items()):
                ws.append([nome, quantidade])
        else:
            ws.append(['Nenhum registro', 0])
        ws.append([])

    @classmethod
    def _nome_pacote(cls, filtros):
        ano = cls._sanitize_filename(str(filtros.get('ano') or datetime.now().year))
        mes = str(filtros.get('mes') or '').strip()
        if mes:
            try:
                mes = f'{int(mes):02d}'
            except (TypeError, ValueError):
                mes = cls._sanitize_filename(mes)
            return f'Documentos_Empresa_{ano}_{mes}_Contador.zip'
        return f'Documentos_Empresa_{ano}_Contador.zip'

    @classmethod
    def _observacoes(cls, documento, metadata, arquivo):
        partes = []
        if not arquivo:
            partes.append('Arquivo nao encontrado ou indisponivel')
        if metadata.get('observacoes'):
            partes.append(str(metadata.get('observacoes')))
        if documento.observacoes:
            partes.append(str(documento.observacoes))
        return '; '.join(partes)

    @staticmethod
    def _entidade_vinculada(vinculo):
        if not vinculo:
            return ''
        if vinculo.resumo_entidade:
            return vinculo.resumo_entidade
        if vinculo.entidade_id:
            return f'{vinculo.tipo_entidade} #{vinculo.entidade_id}'
        return vinculo.tipo_entidade or ''

    @staticmethod
    def _data(valor):
        if not valor:
            return None
        if hasattr(valor, 'strftime'):
            return valor.strftime('%d/%m/%Y')
        texto = str(valor)
        try:
            return datetime.strptime(texto[:10], '%Y-%m-%d').strftime('%d/%m/%Y')
        except ValueError:
            return texto

    @staticmethod
    def _descricao_filtros(filtros):
        pares = []
        for chave in sorted(filtros.keys()):
            valor = filtros.get(chave)
            if valor not in {None, ''}:
                pares.append(f'{chave}={valor}')
        return ', '.join(pares) if pares else 'Sem filtros adicionais'

    @staticmethod
    def _sanitize_filename(valor, limite=80):
        texto = unicodedata.normalize('NFKD', str(valor or 'Documento'))
        texto = texto.encode('ascii', 'ignore').decode('ascii')
        texto = re.sub(r'[\/\\:*?"<>|]+', ' ', texto)
        texto = re.sub(r'\s+', '_', texto).strip('._ ')
        texto = re.sub(r'[^A-Za-z0-9._-]+', '', texto)
        return (texto or 'Documento')[:limite]

    @staticmethod
    def _extensao_arquivo(arquivo):
        if not arquivo:
            return '.bin'
        nome = arquivo.nome_arquivo or ''
        match = re.search(r'(\.[A-Za-z0-9]{1,8})$', nome)
        if match:
            return match.group(1).lower()
        extensao = mimetypes.guess_extension(arquivo.mime_type or '')
        if extensao == '.jpe':
            return '.jpg'
        return extensao or '.bin'

    @staticmethod
    def _parse_int(valor):
        if valor in {None, ''}:
            return None
        try:
            return int(valor)
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _formatar_planilha(ws):
        header_fill = PatternFill('solid', fgColor='1f2937')
        header_font = Font(color='FFFFFF', bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for column_cells in ws.columns:
            width = min(max(len(str(cell.value or '')) for cell in column_cells) + 2, 48)
            ws.column_dimensions[column_cells[0].column_letter].width = width
