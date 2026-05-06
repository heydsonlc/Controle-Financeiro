from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

try:
    from backend.services.ir_relatorio_service import IrRelatorioService
    from backend.services.lastro_financeiro_service import LastroFinanceiroService
    from backend.services.perfil_financeiro_service import PerfilFinanceiroService
except ImportError:
    from services.ir_relatorio_service import IrRelatorioService
    from services.lastro_financeiro_service import LastroFinanceiroService
    from services.perfil_financeiro_service import PerfilFinanceiroService


class LastroRelatorioService:
    MOEDA_FORMATO = 'R$ #,##0.00'
    PENDENCIAS_CONTADOR = {'AGUARDANDO_CONTADOR', 'PENDENTE', 'DIVERGENTE'}
    ERRO_PERFIL = 'Relatorio de lastro empresarial disponivel apenas no perfil Empresa.'

    @classmethod
    def gerar_excel(cls, filtros=None):
        contexto = cls._contexto_empresa()
        filtros = cls._validar_filtros(filtros or {})
        saidas = LastroFinanceiroService.listar_saidas_sem_documento(filtros)
        resumo = cls._resumo(saidas)

        wb = Workbook()
        wb.remove(wb.active)
        cls._aba_resumo(wb, contexto, filtros, resumo)
        cls._aba_saidas(wb, saidas)
        cls._aba_pendencias_contador(wb, saidas)

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer, cls._nome_arquivo(contexto, filtros, 'xlsx')

    @classmethod
    def gerar_pdf(cls, filtros=None):
        contexto = cls._contexto_empresa()
        filtros = cls._validar_filtros(filtros or {})
        saidas = LastroFinanceiroService.listar_saidas_sem_documento(filtros)
        resumo = cls._resumo(saidas)
        linhas = cls._linhas_pdf(contexto, filtros, saidas, resumo)
        return BytesIO(IrRelatorioService._pdf_simples(linhas)), cls._nome_arquivo(contexto, filtros, 'pdf')

    @classmethod
    def _contexto_empresa(cls):
        perfil_id = PerfilFinanceiroService.obter_perfil_ativo_id()
        perfil = PerfilFinanceiroService.obter_perfil_por_id(perfil_id)
        if str(getattr(perfil, 'tipo', '') or '').upper() != 'EMPRESA':
            raise PermissionError(cls.ERRO_PERFIL)
        return {
            'perfil': perfil,
            'gerado_em': datetime.now(),
        }

    @classmethod
    def _validar_filtros(cls, filtros):
        filtros = dict(filtros or {})
        ano = filtros.get('ano')
        if ano not in {None, ''}:
            try:
                int(ano)
            except (TypeError, ValueError) as exc:
                raise ValueError('Ano do relatorio invalido') from exc
        mes = filtros.get('mes')
        if mes not in {None, ''}:
            try:
                mes_int = int(mes)
            except (TypeError, ValueError) as exc:
                raise ValueError('Mes do relatorio invalido') from exc
            if mes_int < 1 or mes_int > 12:
                raise ValueError('Mes do relatorio invalido')
        return filtros

    @classmethod
    def _resumo(cls, saidas):
        total = Decimal('0')
        por_origem = defaultdict(lambda: {'quantidade': 0, 'valor': Decimal('0')})
        por_natureza = defaultdict(lambda: {'quantidade': 0, 'valor': Decimal('0')})
        por_status = defaultdict(lambda: {'quantidade': 0, 'valor': Decimal('0')})
        status_contagem = defaultdict(int)

        for saida in saidas:
            valor = cls._decimal(saida.get('valor'))
            status = str(saida.get('status_lastro') or 'SEM_DOCUMENTO').upper()
            natureza = str(saida.get('natureza_sugerida') or 'OUTRO').upper()
            origem = saida.get('origem') or 'Outros'
            total += valor
            cls._somar_grupo(por_origem, origem, valor)
            cls._somar_grupo(por_natureza, natureza, valor)
            cls._somar_grupo(por_status, status, valor)
            status_contagem[status] += 1

        return {
            'quantidade': len(saidas),
            'valor_total': total,
            'aguardando_contador': status_contagem['AGUARDANDO_CONTADOR'],
            'nao_aplicavel': status_contagem['NAO_APLICAVEL'],
            'pendente': status_contagem['PENDENTE'] + status_contagem['SEM_DOCUMENTO'],
            'por_origem': cls._ordenar_grupos(por_origem),
            'por_natureza': cls._ordenar_grupos(por_natureza),
            'por_status': cls._ordenar_grupos(por_status),
        }

    @classmethod
    def _aba_resumo(cls, wb, contexto, filtros, resumo):
        ws = wb.create_sheet('Resumo')
        linhas = [
            ('Perfil financeiro', contexto['perfil'].nome),
            ('Modo', 'Documentos Fiscais e Lastro'),
            ('Ano', filtros.get('ano') or 'Todos'),
            ('Mes', filtros.get('mes') or 'Todos'),
            ('Data/hora de geracao', contexto['gerado_em'].strftime('%d/%m/%Y %H:%M')),
            ('Total de saidas sem documento', resumo['quantidade']),
            ('Valor total sem documento', float(resumo['valor_total'])),
            ('Quantidade aguardando contador', resumo['aguardando_contador']),
            ('Quantidade nao aplicavel', resumo['nao_aplicavel']),
            ('Quantidade pendente', resumo['pendente']),
        ]
        for linha in linhas:
            ws.append(linha)
        ws.append([])
        cls._adicionar_grupo(ws, 'Total por origem', resumo['por_origem'])
        cls._adicionar_grupo(ws, 'Total por natureza', resumo['por_natureza'])
        cls._adicionar_grupo(ws, 'Total por status', resumo['por_status'])
        cls._formatar_planilha(ws)

    @classmethod
    def _adicionar_grupo(cls, ws, titulo, dados):
        ws.append([titulo])
        ws.append(['Nome', 'Quantidade', 'Valor total'])
        if dados:
            for item in dados:
                ws.append([item['nome'], item['quantidade'], float(item['valor'])])
        else:
            ws.append(['Nenhum registro encontrado', 0, 0.0])
        ws.append([])

    @classmethod
    def _aba_saidas(cls, wb, saidas):
        ws = wb.create_sheet('Saidas_sem_documento')
        ws.append([
            'Data',
            'Origem',
            'Descricao',
            'Categoria',
            'Valor',
            'Natureza sugerida',
            'Status do lastro',
            'Status financeiro',
            'Tipo da entidade',
            'ID da entidade',
            'Observacoes',
        ])
        for saida in saidas:
            ws.append(cls._linha_saida(saida))
        if not saidas:
            ws.append(['Nenhuma saida sem documento encontrada para os filtros selecionados.'])
        cls._formatar_planilha(ws)

    @classmethod
    def _aba_pendencias_contador(cls, wb, saidas):
        ws = wb.create_sheet('Pendencias_contador')
        ws.append(['Data', 'Descricao', 'Valor', 'Natureza', 'Status', 'Observacoes', 'Origem'])
        pendencias = [
            saida for saida in saidas
            if str(saida.get('status_lastro') or '').upper() in cls.PENDENCIAS_CONTADOR
        ]
        for saida in pendencias:
            ws.append([
                cls._data(saida.get('data')),
                saida.get('descricao'),
                float(cls._decimal(saida.get('valor'))),
                saida.get('natureza_sugerida'),
                saida.get('status_lastro'),
                saida.get('observacoes'),
                saida.get('origem'),
            ])
        if not pendencias:
            ws.append(['Nenhuma pendencia para contador encontrada para os filtros selecionados.'])
        cls._formatar_planilha(ws)

    @classmethod
    def _linha_saida(cls, saida):
        return [
            cls._data(saida.get('data')),
            saida.get('origem'),
            saida.get('descricao'),
            saida.get('categoria'),
            float(cls._decimal(saida.get('valor'))),
            saida.get('natureza_sugerida'),
            saida.get('status_lastro'),
            saida.get('status_pagamento'),
            saida.get('tipo_entidade'),
            saida.get('entidade_id'),
            saida.get('observacoes'),
        ]

    @classmethod
    def _linhas_pdf(cls, contexto, filtros, saidas, resumo):
        linhas = [
            'Relatorio de Saidas sem Documento',
            f"Perfil financeiro: {contexto['perfil'].nome}",
            f"Periodo: {cls._periodo_legivel(filtros)}",
            f"Gerado em: {contexto['gerado_em'].strftime('%d/%m/%Y %H:%M')}",
            '',
            'Resumo executivo',
            f"Saidas sem documento: {resumo['quantidade']}",
            f"Valor total sem lastro: {IrRelatorioService._moeda(resumo['valor_total'])}",
            f"Aguardando contador: {resumo['aguardando_contador']}",
            f"Nao aplicavel: {resumo['nao_aplicavel']}",
            f"Pendente: {resumo['pendente']}",
            '',
            'Totais por origem:',
        ]
        linhas.extend(cls._linhas_grupo_pdf(resumo['por_origem']))
        linhas.extend(['', 'Totais por natureza:'])
        linhas.extend(cls._linhas_grupo_pdf(resumo['por_natureza']))
        linhas.extend(['', 'Totais por status:'])
        linhas.extend(cls._linhas_grupo_pdf(resumo['por_status']))
        linhas.extend(['', 'Saidas:'])
        if saidas:
            for saida in saidas[:90]:
                linhas.append(' | '.join([
                    cls._data(saida.get('data')) or '-',
                    str(saida.get('origem') or '-'),
                    cls._texto(saida.get('descricao'), 34),
                    IrRelatorioService._moeda(saida.get('valor') or 0),
                    str(saida.get('natureza_sugerida') or '-'),
                    str(saida.get('status_lastro') or '-'),
                ]))
        else:
            linhas.append('Nenhuma saida sem documento encontrada para os filtros selecionados.')
        linhas.extend([
            '',
            'Relatorio gerencial. A classificacao fiscal/contabil deve ser validada pelo contador responsavel.',
        ])
        return linhas

    @staticmethod
    def _linhas_grupo_pdf(dados):
        if not dados:
            return ['- Nenhum registro encontrado.']
        return [
            f"- {item['nome']}: {item['quantidade']} item(ns), {IrRelatorioService._moeda(item['valor'])}"
            for item in dados
        ]

    @staticmethod
    def _periodo_legivel(filtros):
        ano = filtros.get('ano') or 'Todos'
        mes = filtros.get('mes')
        return f'{mes}/{ano}' if mes else str(ano)

    @staticmethod
    def _nome_arquivo(contexto, filtros, extensao):
        ano = filtros.get('ano') or datetime.now().year
        mes = filtros.get('mes')
        perfil = re.sub(r'[^A-Za-z0-9_-]+', '_', contexto['perfil'].nome or 'Empresa').strip('_')
        periodo = f'{ano}_{int(mes):02d}' if mes else str(ano)
        return f'Saidas_sem_documento_{perfil}_{periodo}.{extensao}'

    @staticmethod
    def _formatar_planilha(ws):
        header_fill = PatternFill('solid', fgColor='1d4ed8')
        header_font = Font(color='FFFFFF', bold=True)
        header_labels = {
            'Data',
            'Perfil financeiro',
            'Nome',
            'Total por origem',
            'Total por natureza',
            'Total por status',
        }
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row == 1 or cell.value in header_labels:
                    cell.fill = header_fill
                    cell.font = header_font
                if isinstance(cell.value, float):
                    cell.number_format = LastroRelatorioService.MOEDA_FORMATO
        ws.freeze_panes = 'A2'
        if ws.max_row > 1 and ws.max_column > 1:
            ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            letra = column[0].column_letter
            tamanho = max(len(str(cell.value or '')) for cell in column)
            ws.column_dimensions[letra].width = min(max(tamanho + 2, 12), 48)

    @staticmethod
    def _somar_grupo(grupos, chave, valor):
        grupos[chave]['quantidade'] += 1
        grupos[chave]['valor'] += valor

    @staticmethod
    def _ordenar_grupos(grupos):
        return [
            {'nome': chave, 'quantidade': dados['quantidade'], 'valor': dados['valor']}
            for chave, dados in sorted(grupos.items(), key=lambda item: (-item[1]['valor'], item[0]))
        ]

    @staticmethod
    def _data(valor):
        if not valor:
            return None
        if hasattr(valor, 'strftime'):
            return valor.strftime('%d/%m/%Y')
        try:
            return datetime.fromisoformat(str(valor)).strftime('%d/%m/%Y')
        except ValueError:
            return str(valor)

    @staticmethod
    def _decimal(valor):
        try:
            return Decimal(str(valor or 0))
        except (InvalidOperation, TypeError, ValueError):
            return Decimal('0')

    @staticmethod
    def _texto(valor, limite):
        texto = str(valor or '').strip()
        return texto[:limite] if texto else '-'
