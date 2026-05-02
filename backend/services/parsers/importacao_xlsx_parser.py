from datetime import date, datetime

try:
    from backend.services.parsers.importacao_csv_parser import converter_tabela_para_payload
except ImportError:
    from services.parsers.importacao_csv_parser import converter_tabela_para_payload


def _valor_para_texto(valor):
    if valor is None:
        return ''
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    return str(valor).strip()


def _primeira_tabela(sheet):
    linhas = []
    for row in sheet.iter_rows(values_only=True):
        valores = [_valor_para_texto(celula) for celula in row]
        if any(valores):
            linhas.append(valores)

    if not linhas:
        raise ValueError('XLSX vazio ou sem dados reconheciveis.')

    cabecalho = linhas[0]
    dados = [linha for linha in linhas[1:] if any(str(c).strip() for c in linha)]
    if not dados:
        raise ValueError('XLSX sem linhas de lancamentos.')
    return cabecalho, dados


def parse(file_storage, competencia=None):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError('Dependencia openpyxl nao instalada. Execute pip install -r requirements.txt.') from exc

    if hasattr(file_storage, 'seek'):
        file_storage.seek(0)
    elif hasattr(file_storage, 'stream') and hasattr(file_storage.stream, 'seek'):
        file_storage.stream.seek(0)

    origem = getattr(file_storage, 'stream', file_storage)
    workbook = load_workbook(origem, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        colunas, linhas_dados = _primeira_tabela(sheet)
    finally:
        workbook.close()

    payload = converter_tabela_para_payload(
        colunas=colunas,
        linhas_dados=linhas_dados,
        origem='xlsx',
        arquivo_nome=getattr(file_storage, 'filename', None),
    )
    if payload.get('requer_mapeamento'):
        raise ValueError('XLSX sem cabecalhos reconheciveis. Use colunas equivalentes a CSV Nubank ou Caixa.')
    return payload
