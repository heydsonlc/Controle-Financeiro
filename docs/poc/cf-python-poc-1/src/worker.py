"""CF-PYTHON-POC-1 — diagnostico isolado do Controle Financeiro em Python Workers.

Sem credenciais reais e sem banco real. Cada rota testa uma camada e devolve
JSON com o resultado, para que uma falha numa camada nao esconda as outras.
"""
import importlib
import io
import os
import sys
import time
import traceback

from flask import Flask, jsonify
from workers import wsgi

app = Flask(__name__)


def _erro(exc):
    return f"{type(exc).__name__}: {exc}"[:400]


def _tentar(func):
    inicio = time.perf_counter()
    try:
        resultado = func()
        return {"ok": True, "resultado": resultado, "ms": round((time.perf_counter() - inicio) * 1000, 1)}
    except BaseException as exc:  # noqa: BLE001 - diagnostico
        return {"ok": False, "erro": _erro(exc), "ms": round((time.perf_counter() - inicio) * 1000, 1),
                "trace": traceback.format_exc(limit=3)[-600:]}


@app.get("/")
def index():
    return {"ok": True, "python": sys.version, "platform": sys.platform}


@app.get("/deps")
def deps():
    modulos = [
        "flask", "flask_sqlalchemy", "sqlalchemy", "flask_login", "flask_cors", "flask_migrate",
        "flask_wtf", "dotenv", "dateutil", "apscheduler", "openpyxl", "pdfminer", "pdfplumber",
        "PIL", "pypdfium2", "cryptography", "pg8000", "psycopg", "psycopg2", "sqlite3",
    ]
    return jsonify({m: _tentar(lambda m=m: getattr(importlib.import_module(m), "__version__", "importado"))
                    for m in modulos})


def _pdf_sintetico(texto="Documento de teste POC"):
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        None,
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 20 100 Td ({texto}) Tj ET".encode()
    objetos[3] = b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
    saida = io.BytesIO()
    saida.write(b"%PDF-1.4\n")
    offsets = []
    for i, corpo in enumerate(objetos, start=1):
        offsets.append(saida.tell())
        saida.write(f"{i} 0 obj\n".encode() + corpo + b"\nendobj\n")
    xref = saida.tell()
    saida.write(f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        saida.write(f"{off:010d} 00000 n \n".encode())
    saida.write(f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    return saida.getvalue()


@app.get("/pdf")
def pdf():
    dados = _pdf_sintetico()

    def extrair():
        import pdfplumber
        with pdfplumber.open(io.BytesIO(dados)) as doc:
            return doc.pages[0].extract_text()

    def renderizar():
        import pypdfium2
        doc = pypdfium2.PdfDocument(dados)
        imagem = doc[0].render(scale=1).to_pil()
        return list(imagem.size)

    def to_image_pdfplumber():
        import pdfplumber
        with pdfplumber.open(io.BytesIO(dados)) as doc:
            return list(doc.pages[0].to_image(resolution=36).original.size)

    return jsonify({"extract_text": _tentar(extrair), "pypdfium2_render": _tentar(renderizar),
                    "pdfplumber_to_image": _tentar(to_image_pdfplumber)})


@app.get("/xlsx")
def xlsx():
    def gerar():
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(["competencia", "valor"])
        wb.active.append(["2026-10", 123.45])
        buf = io.BytesIO()
        wb.save(buf)
        return len(buf.getvalue())
    return jsonify({"openpyxl_bytesio": _tentar(gerar)})


@app.get("/runtime")
def runtime():
    def fs_escrita():
        caminho = "/tmp/poc_contador.txt"
        atual = int(open(caminho).read()) if os.path.exists(caminho) else 0
        with open(caminho, "w") as f:
            f.write(str(atual + 1))
        return {"contador_neste_isolate": atual + 1, "cwd": os.getcwd()}

    def thread():
        import threading
        marcador = []
        t = threading.Thread(target=lambda: marcador.append(1))
        t.start()
        t.join(timeout=1)
        return {"executou": bool(marcador)}

    def subprocesso():
        import subprocess
        return subprocess.run(["echo", "oi"], capture_output=True, text=True, timeout=5).stdout

    def tempfile_():
        import tempfile
        with tempfile.NamedTemporaryFile(delete=True) as tmp:
            tmp.write(b"x")
            return tmp.name

    def memoria_wasm():
        import pyodide_js  # noqa: F401
        return {"heap_bytes": int(pyodide_js._module.HEAP8.length)}

    return jsonify({"fs_escrita": _tentar(fs_escrita), "thread": _tentar(thread),
                    "subprocess": _tentar(subprocesso), "tempfile": _tentar(tempfile_),
                    "memoria_wasm": _tentar(memoria_wasm)})


@app.get("/db-driver")
def db_driver():
    """Sem banco real: so engine/dialeto e uma tentativa a uma porta local fechada."""
    def engine_pg8000():
        from sqlalchemy import create_engine
        eng = create_engine("postgresql+pg8000://poc:poc@127.0.0.1:1/poc")
        return {"dialeto": eng.dialect.name, "driver": eng.dialect.driver}

    def conectar_porta_fechada():
        import pg8000.dbapi
        pg8000.dbapi.connect(user="poc", password="poc", host="127.0.0.1", port=1, database="poc", timeout=3)
        return "conectou (inesperado)"

    def engine_psycopg3():
        from sqlalchemy import create_engine
        eng = create_engine("postgresql+psycopg://poc:poc@127.0.0.1:1/poc")
        return {"dialeto": eng.dialect.name, "driver": eng.dialect.driver}

    return jsonify({"sqlalchemy_pg8000": _tentar(engine_pg8000),
                    "pg8000_socket_porta_fechada": _tentar(conectar_porta_fechada),
                    "sqlalchemy_psycopg3": _tentar(engine_psycopg3)})


HASH_HOST_SCRYPT = (
    "scrypt:32768:8:1$C2TPKandIiXtlj5p$3e80f290f2a21333da9a179a8c58e989c71350932b9ec4587da8a6ad51ee11ba"
    "0abbed855d893220e0616e10cfd0356ef6abf5b9365959647d409685488efb39"
)


def _instalar_shim_scrypt():
    """Ajuste necessario (e testado aqui): hashlib.scrypt via `cryptography`."""
    import hashlib
    if hasattr(hashlib, "scrypt"):
        return "nativo"

    def scrypt(password, *, salt, n, r, p, maxmem=0, dklen=64):
        from cryptography.hazmat.primitives.kdf.scrypt import Scrypt
        return Scrypt(salt=salt, length=dklen, n=n, r=r, p=p).derive(password)

    hashlib.scrypt = scrypt
    return "shim cryptography"


@app.get("/kdf")
def kdf():
    import hashlib
    resultado = {"hashlib_tem_scrypt": hasattr(hashlib, "scrypt"),
                 "hashlib_tem_pbkdf2_hmac": hasattr(hashlib, "pbkdf2_hmac")}

    def verificar_hash_do_host():
        _instalar_shim_scrypt()
        from werkzeug.security import check_password_hash
        return {"senha_certa": check_password_hash(HASH_HOST_SCRYPT, "senha-poc-123456"),
                "senha_errada": check_password_hash(HASH_HOST_SCRYPT, "outra")}

    resultado["verifica_hash_scrypt_gerado_no_cpython"] = _tentar(verificar_hash_do_host)
    return jsonify(resultado)


_APP_REAL = {}


@app.get("/app-real")
def app_real():
    """Executa o codigo real do Controle Financeiro (commit HEAD), sem adaptacoes,
    com config 'testing' (SQLite em memoria). Nada de banco real."""
    resultado = {}

    def importar_e_criar():
        os.environ.setdefault("APP_ENV", "testing")
        _instalar_shim_scrypt()
        inicio = time.perf_counter()
        from backend.app import create_app  # importa tambem app = create_app() no modulo
        from backend.models import Usuario, db
        real = create_app("testing")
        with real.app_context():
            db.create_all()
            u = Usuario(email="poc@local.invalid", nome="POC", ativo=True)
            # Hash gerado no CPython do host (Werkzeug, scrypt padrao) para senha ficticia:
            # simula um usuario JA existente no banco. Verificado aqui via shim de scrypt.
            u.password_hash = HASH_HOST_SCRYPT
            db.session.add(u)
            db.session.commit()
        _APP_REAL["app"] = real
        return {"tabelas": len(db.metadata.tables), "ms_import_create": round((time.perf_counter() - inicio) * 1000)}

    resultado["import_create_app"] = _tentar(importar_e_criar)
    real = _APP_REAL.get("app")
    if real is None:
        return jsonify(resultado)

    client = real.test_client()

    def req(metodo, rota, **kw):
        r = getattr(client, metodo)(rota, **kw)
        return {"status": r.status_code, "location": r.headers.get("Location"),
                "tipo": r.content_type, "bytes": len(r.data)}

    resultado["GET /health"] = _tentar(lambda: req("get", "/health"))
    resultado["GET / sem sessao"] = _tentar(lambda: req("get", "/"))
    resultado["GET /login (Jinja)"] = _tentar(lambda: req("get", "/login"))
    resultado["POST /login"] = _tentar(lambda: req("post", "/login",
                                                   data={"email": "poc@local.invalid", "senha": "senha-poc-123456"}))
    resultado["GET / com sessao (dashboard)"] = _tentar(lambda: req("get", "/"))
    for rota in ["/api/categorias/", "/api/despesas/", "/api/dashboard/resumo", "/api/contas/",
                 "/api/financiamentos/", "/api/cartoes/", "/api/receitas/", "/api/backup/status",
                 "/api/ocr/status", "/api/ir/comprovantes"]:
        resultado[f"GET {rota}"] = _tentar(lambda rota=rota: req("get", rota))
    return jsonify(resultado)


Default = wsgi.entrypoint(app)
