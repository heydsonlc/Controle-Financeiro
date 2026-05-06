from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.app import create_app
from backend.models import (
    Categoria,
    DespesaPrevista,
    IrComprovante,
    IrComprovanteArquivo,
    LastroFinanceiroPendencia,
    db,
)
from backend.services.perfil_financeiro_service import PerfilFinanceiroService


@pytest.fixture()
def app():
    app = create_app('testing')
    with app.app_context():
        db.create_all()
        PerfilFinanceiroService.obter_ou_criar_perfis_iniciais()
        db.session.commit()
        yield app
        db.session.remove()
        db.drop_all()


@pytest.fixture()
def client(app):
    return app.test_client()


def _perfis(client):
    return client.get('/api/perfis-financeiros').get_json()['perfis']


def _perfil(client, nome):
    return next(perfil for perfil in _perfis(client) if perfil['nome'] == nome)


def _trocar_perfil(client, nome):
    perfil = _perfil(client, nome)
    response = client.post('/api/perfis-financeiros/ativo', json={'perfil_id': perfil['id']})
    assert response.status_code == 200, response.get_data(as_text=True)
    return perfil


def _categoria(perfil_id, nome='Operacional'):
    categoria = Categoria(
        perfil_financeiro_id=perfil_id,
        nome=nome,
        descricao='Categoria para relatorio de lastro',
        ativo=True,
    )
    db.session.add(categoria)
    db.session.flush()
    return categoria


def _saida(perfil_id, nome='Despesa operacional', valor='250.00', categoria=None):
    if categoria is None:
        nome_categoria = nome
        if Categoria.query.filter_by(perfil_financeiro_id=perfil_id, nome=nome_categoria).first():
            nome_categoria = f'{nome} {valor}'
        categoria = _categoria(perfil_id, nome_categoria)
    despesa = DespesaPrevista(
        perfil_financeiro_id=perfil_id,
        origem_tipo='MANUAL',
        origem_id=1,
        categoria_id=categoria.id,
        data_prevista=date(2026, 5, 10),
        data_original_prevista=date(2026, 5, 10),
        data_atual_prevista=date(2026, 5, 10),
        valor_previsto=Decimal(valor),
        status='PREVISTA',
    )
    db.session.add(despesa)
    db.session.flush()
    return despesa


def _pendencia(perfil_id, saida, status='AGUARDANDO_CONTADOR', natureza='PATRIMONIO_IMOBILIZADO', observacoes='Validar com contador'):
    pendencia = LastroFinanceiroPendencia(
        perfil_financeiro_id=perfil_id,
        tipo_entidade='DESPESA_PREVISTA',
        entidade_id=saida.id,
        status_lastro=status,
        natureza=natureza,
        observacoes=observacoes,
        ativo=True,
    )
    db.session.add(pendencia)
    db.session.flush()
    return pendencia


def _comprovante_com_binario(perfil_id):
    comprovante = IrComprovante(
        perfil_financeiro_id=perfil_id,
        ano_calendario=2026,
        data_documento=date(2026, 5, 11),
        prestador_nome='Fornecedor com arquivo',
        valor=Decimal('99.00'),
        status='VALIDADO',
        hash_arquivo=f'hash-lastro-relatorio-{perfil_id}',
    )
    db.session.add(comprovante)
    db.session.flush()
    arquivo = IrComprovanteArquivo(
        comprovante_id=comprovante.id,
        nome_arquivo='nota-secreta.pdf',
        mime_type='application/pdf',
        tamanho_bytes=24,
        conteudo=b'%PDF-1.4 SEGREDO_BINARIO',
        hash_arquivo=comprovante.hash_arquivo,
    )
    db.session.add(arquivo)
    db.session.flush()
    return comprovante


def _linhas_sheet(ws):
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def test_excel_saidas_sem_documento_empresa_tem_abas(client, app):
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        _saida(empresa['id'], valor='300.00')
        db.session.commit()

    response = client.get('/api/ir/lastro/saidas-sem-documento/relatorio-excel?ano=2026&mes=5')
    wb = load_workbook(BytesIO(response.data))

    assert response.status_code == 200
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert {'Resumo', 'Saidas_sem_documento', 'Pendencias_contador'}.issubset(set(wb.sheetnames))
    assert wb['Saidas_sem_documento']['E2'].value == 300


def test_pdf_saidas_sem_documento_empresa_e_pdf_real(client, app):
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        _saida(empresa['id'], nome='Notebook empresa', valor='1200.00')
        db.session.commit()

    response = client.get('/api/ir/lastro/saidas-sem-documento/relatorio-pdf?ano=2026&mes=5')

    assert response.status_code == 200
    assert response.mimetype == 'application/pdf'
    assert response.data.startswith(b'%PDF')
    assert b'Relatorio de Saidas sem Documento' in response.data
    assert b'Notebook empresa' in response.data


def test_perfil_pessoal_recebe_erro_controlado(client):
    response = client.get('/api/ir/lastro/saidas-sem-documento/relatorio-excel?ano=2026')

    assert response.status_code == 403
    assert response.get_json()['success'] is False
    assert 'perfil Empresa' in response.get_json()['error']


def test_relatorio_respeita_filtros_de_status_e_natureza(client, app):
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        pendente = _saida(empresa['id'], nome='Despesa operacional', valor='100.00')
        contador = _saida(empresa['id'], nome='Notebook patrimonio', valor='900.00')
        _pendencia(empresa['id'], contador, status='AGUARDANDO_CONTADOR', natureza='PATRIMONIO_IMOBILIZADO')
        db.session.commit()
        pendente_id = pendente.id
        contador_id = contador.id

    response = client.get(
        '/api/ir/lastro/saidas-sem-documento/relatorio-excel'
        '?ano=2026&mes=5&status=AGUARDANDO_CONTADOR&natureza=PATRIMONIO_IMOBILIZADO'
    )
    wb = load_workbook(BytesIO(response.data))
    ids = [row[9].value for row in wb['Saidas_sem_documento'].iter_rows(min_row=2)]

    assert response.status_code == 200
    assert contador_id in ids
    assert pendente_id not in ids
    assert wb['Resumo']['B6'].value == 1
    assert wb['Resumo']['B7'].value == 900


def test_relatorio_nao_mistura_outro_perfil(client, app):
    pessoal = _perfil(client, 'Pessoal')
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        _saida(pessoal['id'], nome='Despesa pessoal oculta', valor='500.00')
        _saida(empresa['id'], nome='Despesa empresa visivel', valor='250.00')
        db.session.commit()

    response = client.get('/api/ir/lastro/saidas-sem-documento/relatorio-excel?ano=2026&mes=5')
    wb = load_workbook(BytesIO(response.data))
    descricoes = [row[2] for row in _linhas_sheet(wb['Saidas_sem_documento'])]

    assert 'Despesa empresa visivel - PREVISTA' in descricoes
    assert 'Despesa pessoal oculta - PREVISTA' not in descricoes


def test_totais_quantidade_e_valor_estao_corretos(client, app):
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        _saida(empresa['id'], valor='100.50')
        _saida(empresa['id'], valor='200.25')
        db.session.commit()

    response = client.get('/api/ir/lastro/saidas-sem-documento/relatorio-excel?ano=2026&mes=5')
    wb = load_workbook(BytesIO(response.data))

    assert wb['Resumo']['B6'].value == 2
    assert wb['Resumo']['B7'].value == 300.75


def test_relatorio_vazio_funciona(client):
    _trocar_perfil(client, 'Empresa')

    response = client.get('/api/ir/lastro/saidas-sem-documento/relatorio-excel?ano=2026&mes=5')
    wb = load_workbook(BytesIO(response.data))

    assert response.status_code == 200
    assert wb['Saidas_sem_documento']['A2'].value == 'Nenhuma saida sem documento encontrada para os filtros selecionados.'
    assert wb['Pendencias_contador']['A2'].value == 'Nenhuma pendencia para contador encontrada para os filtros selecionados.'


def test_relatorio_nao_expoe_texto_extraido_ou_binario(client, app):
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        _saida(empresa['id'], valor='100.00')
        _comprovante_com_binario(empresa['id'])
        db.session.commit()

    excel = client.get('/api/ir/lastro/saidas-sem-documento/relatorio-excel?ano=2026&mes=5')
    pdf = client.get('/api/ir/lastro/saidas-sem-documento/relatorio-pdf?ano=2026&mes=5')

    assert b'SEGREDO_BINARIO' not in excel.data
    assert b'SEGREDO_BINARIO' not in pdf.data
    assert b'%PDF-1.4 SEGREDO_BINARIO' not in pdf.data


def test_botoes_do_frontend_existentes():
    template = (Path(__file__).resolve().parents[1] / 'frontend' / 'templates' / 'imposto_renda.html').read_text(encoding='utf-8')
    js = (Path(__file__).resolve().parents[1] / 'frontend' / 'static' / 'js' / 'imposto_renda.js').read_text(encoding='utf-8')

    assert 'ir-saidas-btn-excel' in template
    assert 'ir-saidas-btn-pdf' in template
    assert 'relatorio-excel' in js
    assert 'relatorio-pdf' in js


def test_sem_url_externa_ou_dependencia_nova():
    service = (Path(__file__).resolve().parents[1] / 'backend' / 'services' / 'lastro_relatorio_service.py').read_text(encoding='utf-8')
    template = (Path(__file__).resolve().parents[1] / 'frontend' / 'templates' / 'imposto_renda.html').read_text(encoding='utf-8')

    assert 'http://' not in service
    assert 'https://' not in service
    assert 'http://' not in template
    assert 'https://' not in template
    assert 'reportlab' not in service.lower()
    assert 'weasyprint' not in service.lower()
