from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from backend.app import create_app
from backend.models import (
    Categoria,
    DespesaPrevista,
    IrCategoria,
    IrComprovante,
    IrComprovanteArquivo,
    IrComprovanteVinculo,
    PerfilFinanceiro,
    db,
)
from backend.services.perfil_financeiro_service import PerfilFinanceiroService
from tests.conftest import autenticar_cliente_teste


@pytest.fixture()
def app():
    app = create_app('testing')
    with app.app_context():
        db.create_all()
        PerfilFinanceiroService.obter_ou_criar_perfis_iniciais()
        db.session.commit()
        yield app
        db.session.remove()
        db.drop_all()


@pytest.fixture()
def client(app):
    return autenticar_cliente_teste(app.test_client(), app)


def _perfis(client):
    return client.get('/api/perfis-financeiros').get_json()['perfis']


def _perfil(client, nome):
    return next(perfil for perfil in _perfis(client) if perfil['nome'] == nome)


def _trocar_perfil(client, nome):
    perfil = _perfil(client, nome)
    response = client.post('/api/perfis-financeiros/ativo', json={'perfil_id': perfil['id']})
    assert response.status_code == 200, response.get_data(as_text=True)
    return perfil


def _criar_comprovante(app, perfil_id, nome, categoria_ir_nome='Saude', valor='100.00', status='VALIDADO', tipo_contexto='PESSOAL'):
    with app.app_context():
        categoria = Categoria(
            perfil_financeiro_id=perfil_id,
            nome=f'Categoria {nome}',
            descricao='Teste relatorio',
            ativo=True,
        )
        categoria_ir = IrCategoria.query.filter_by(nome=categoria_ir_nome).first()
        if not categoria_ir:
            categoria_ir = IrCategoria(
                nome=categoria_ir_nome,
                descricao='Categoria fiscal teste',
                dedutivel=True,
                ativo=True,
                ordem=1,
                tipo_contexto=tipo_contexto,
                natureza='DESPESA_OPERACIONAL' if tipo_contexto == 'EMPRESA' else 'DEDUCAO_IRPF',
            )
        db.session.add_all([categoria, categoria_ir])
        db.session.flush()
        comprovante = IrComprovante(
            perfil_financeiro_id=perfil_id,
            ano_calendario=2026,
            data_documento=date(2026, 5, 1),
            prestador_nome=nome,
            prestador_cpf_cnpj='01.000.000/0001-00',
            valor=Decimal(valor),
            categoria_id=categoria.id,
            categoria_ir_id=categoria_ir.id,
            dedutivel=True,
            status=status,
            hash_arquivo=f'hash-{perfil_id}-{nome}',
            observacoes='Observacao segura',
        )
        db.session.add(comprovante)
        db.session.flush()
        arquivo = IrComprovanteArquivo(
            comprovante_id=comprovante.id,
            nome_arquivo=f'{nome}.pdf',
            mime_type='application/pdf',
            tamanho_bytes=8,
            conteudo=b'%PDF-1.4',
            hash_arquivo=comprovante.hash_arquivo,
        )
        db.session.add(arquivo)
        db.session.commit()
        return comprovante.id


def _criar_despesa_e_vinculo(app, perfil_id, comprovante_id):
    with app.app_context():
        categoria = Categoria(
            perfil_financeiro_id=perfil_id,
            nome='Despesa vinculavel relatorio',
            descricao='Lastro',
            ativo=True,
        )
        db.session.add(categoria)
        db.session.flush()
        despesa = DespesaPrevista(
            perfil_financeiro_id=perfil_id,
            origem_tipo='MANUAL',
            origem_id=1,
            categoria_id=categoria.id,
            data_prevista=date(2026, 5, 1),
            data_original_prevista=date(2026, 5, 1),
            data_atual_prevista=date(2026, 5, 1),
            valor_previsto=Decimal('250.00'),
            status='PREVISTA',
        )
        db.session.add(despesa)
        db.session.flush()
        vinculo = IrComprovanteVinculo(
            comprovante_id=comprovante_id,
            perfil_financeiro_id=perfil_id,
            tipo_entidade='DESPESA_PREVISTA',
            entidade_id=despesa.id,
            resumo_entidade='Despesa vinculavel relatorio',
            tipo_vinculo='NOTA_FISCAL',
            natureza='DESPESA_OPERACIONAL',
            status_lastro='VALIDADO',
            ativo=True,
        )
        db.session.add(vinculo)
        db.session.commit()


def test_excel_perfil_pessoal_tem_abas_e_dados_do_irpf(client, app):
    pessoal = _perfil(client, 'Pessoal')
    _criar_comprovante(app, pessoal['id'], 'Hospital Pessoal', valor='1250.00')

    response = client.get('/api/ir/relatorios/excel?ano=2026')
    wb = load_workbook(BytesIO(response.data))

    assert response.status_code == 200
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert {'Resumo', 'Comprovantes'}.issubset(set(wb.sheetnames))
    assert wb['Comprovantes']['C2'].value == 'Hospital Pessoal'
    assert 'Lastros' not in wb.sheetnames


def test_pdf_perfil_pessoal_e_pdf_real(client, app):
    pessoal = _perfil(client, 'Pessoal')
    _criar_comprovante(app, pessoal['id'], 'Clinica Pessoal', valor='200.00')

    response = client.get('/api/ir/relatorios/pdf?ano=2026')

    assert response.status_code == 200
    assert response.mimetype == 'application/pdf'
    assert response.data.startswith(b'%PDF')
    assert b'Imposto de Renda' in response.data


def test_excel_empresa_tem_aba_lastros_e_nao_mistura_pessoal(client, app):
    pessoal = _perfil(client, 'Pessoal')
    empresa = _trocar_perfil(client, 'Empresa')
    _criar_comprovante(app, pessoal['id'], 'Hospital Pessoal', valor='100.00')
    comprovante_empresa = _criar_comprovante(app, empresa['id'], 'Fornecedor Empresa', categoria_ir_nome='Despesa operacional', valor='500.00', tipo_contexto='EMPRESA')
    _criar_despesa_e_vinculo(app, empresa['id'], comprovante_empresa)

    response = client.get('/api/ir/relatorios/excel?ano=2026')
    wb = load_workbook(BytesIO(response.data))
    documentos = [row[2].value for row in wb['Documentos'].iter_rows(min_row=2)]

    assert response.status_code == 200
    assert {'Resumo', 'Documentos', 'Lastros'}.issubset(set(wb.sheetnames))
    assert 'Fornecedor Empresa' in documentos
    assert 'Hospital Pessoal' not in documentos
    assert wb['Lastros']['D2'].value == 'DESPESA_PREVISTA'


def test_pdf_empresa_respeita_perfil_ativo(client, app):
    pessoal = _perfil(client, 'Pessoal')
    empresa = _trocar_perfil(client, 'Empresa')
    _criar_comprovante(app, pessoal['id'], 'Documento Pessoal', valor='100.00')
    _criar_comprovante(app, empresa['id'], 'Documento Empresa', categoria_ir_nome='Despesa operacional', valor='300.00', tipo_contexto='EMPRESA')

    response = client.get('/api/ir/relatorios/pdf?ano=2026')

    assert response.status_code == 200
    assert response.data.startswith(b'%PDF')
    assert b'Documentos Fiscais e Lastro' in response.data
    assert b'Documento Empresa' in response.data
    assert b'Documento Pessoal' not in response.data


def test_relatorio_sem_documentos_e_parametro_invalido(client):
    vazio = client.get('/api/ir/relatorios/excel?ano=2026')
    invalido = client.get('/api/ir/relatorios/excel?ano=abc')

    assert vazio.status_code == 200
    assert load_workbook(BytesIO(vazio.data))['Comprovantes'].max_row == 1
    assert invalido.status_code == 400
    assert invalido.get_json()['error'] == 'Ano do relatorio invalido'


def test_relatorio_nao_usa_url_externa_ou_dependencia_pdf_nova():
    service = (Path(__file__).resolve().parents[1] / 'backend' / 'services' / 'ir_relatorio_service.py').read_text(encoding='utf-8')

    assert 'http://' not in service
    assert 'https://' not in service
    assert 'weasyprint' not in service.lower()
    assert 'reportlab' not in service.lower()
