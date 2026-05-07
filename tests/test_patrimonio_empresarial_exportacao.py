from io import BytesIO

from openpyxl import load_workbook
import pytest

from backend.app import create_app
from backend.models import db
from backend.services.perfil_financeiro_service import PerfilFinanceiroService


@pytest.fixture()
def app():
    app = create_app('testing')
    with app.app_context():
        db.create_all()
        PerfilFinanceiroService.obter_ou_criar_perfis_iniciais()
        db.session.commit()
        yield app
        db.session.remove()
        db.drop_all()


@pytest.fixture()
def client(app):
    return app.test_client()


def _perfis(client):
    return client.get('/api/perfis-financeiros').get_json()['perfis']


def _perfil(client, nome):
    return next(perfil for perfil in _perfis(client) if perfil['nome'] == nome)


def _trocar_perfil(client, nome):
    perfil = _perfil(client, nome)
    response = client.post('/api/perfis-financeiros/ativo', json={'perfil_id': perfil['id']})
    assert response.status_code == 200, response.get_data(as_text=True)
    return perfil


def _payload_bem(indice=1, categoria='Informatica', status_documental='SEM_DOCUMENTO'):
    return {
        'nome': f'Bem patrimonial {indice:02d}',
        'codigo': f'TAG-{indice:04d}',
        'categoria': categoria,
        'fornecedor': 'Fornecedor Patrimonial Ltda.',
        'documento_numero': f'DOC-{indice:04d}',
        'data_aquisicao': '2026-05-05',
        'valor_aquisicao': 1000 + indice,
        'vida_util_meses': 60,
        'centro_custo': 'Administrativo',
        'localizacao': 'Escritorio',
        'responsavel': 'Equipe Administrativa',
        'status_documental': status_documental,
    }


def _criar_bens(client, quantidade=3):
    criados = []
    for indice in range(1, quantidade + 1):
        categoria = 'Moveis' if indice % 2 == 0 else 'Informatica'
        payload = _payload_bem(indice, categoria=categoria)
        response = client.post('/api/patrimonio/bens', json=payload)
        assert response.status_code == 201, response.get_data(as_text=True)
        criados.append(response.get_json()['data'])
    return criados


def test_patrimonio_empresa_exporta_excel_com_todos_os_bens_filtrados(client):
    _trocar_perfil(client, 'Empresa')
    _criar_bens(client, 12)

    response = client.get('/api/patrimonio/bens/exportar-excel?categoria=Informatica')

    assert response.status_code == 200, response.get_data(as_text=True)
    assert response.mimetype == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    workbook = load_workbook(BytesIO(response.data))
    assert {'Resumo', 'Bens'} <= set(workbook.sheetnames)
    ws = workbook['Bens']
    headers = [cell.value for cell in ws[1]]
    assert headers[:5] == ['Nome', 'Codigo', 'Categoria', 'Data de aquisicao', 'Valor de aquisicao']
    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(linhas) == 6
    assert all(linha[2] == 'Informatica' for linha in linhas)


def test_patrimonio_empresa_exporta_pdf_com_filtros(client):
    _trocar_perfil(client, 'Empresa')
    _criar_bens(client, 2)

    response = client.get('/api/patrimonio/bens/exportar-pdf?categoria=Moveis')

    assert response.status_code == 200, response.get_data(as_text=True)
    assert response.mimetype == 'application/pdf'
    assert response.data.startswith(b'%PDF')
    assert b'Bem patrimonial 02' in response.data
    assert b'Bem patrimonial 01' not in response.data


def test_patrimonio_exportacao_bloqueia_perfil_pessoal(client):
    _trocar_perfil(client, 'Pessoal')

    excel = client.get('/api/patrimonio/bens/exportar-excel')
    pdf = client.get('/api/patrimonio/bens/exportar-pdf')

    assert excel.status_code == 403
    assert pdf.status_code == 403
    assert 'perfil Empresa' in excel.get_json()['error']


def test_patrimonio_template_tem_acoes_exportacao_filtros_e_paginacao(client):
    html = client.get('/patrimonio').get_data(as_text=True)

    assert 'id="btn-filtros-bens"' in html
    assert 'id="btn-exportar-bens-excel"' in html
    assert 'id="btn-exportar-bens-pdf"' in html
    assert 'documento_relatorio.png' in html
    assert 'Gerar relat&oacute;rio em PDF' in html
    assert 'id="bens-paginacao"' in html
    assert 'patrimonio-empresa-heading' not in html
