from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from backend.app import create_app
from backend.models import (
    DocumentoEmpresarialMetadata,
    IrComprovante,
    IrComprovanteArquivo,
    db,
)
from backend.services.perfil_financeiro_service import PerfilFinanceiroService
from tests.conftest import autenticar_cliente_teste


@pytest.fixture()
def app():
    app = create_app('testing')
    with app.app_context():
        db.create_all()
        PerfilFinanceiroService.obter_ou_criar_perfis_iniciais()
        db.session.commit()
        yield app
        db.session.remove()
        db.drop_all()


@pytest.fixture()
def client(app):
    return autenticar_cliente_teste(app.test_client(), app)


def _perfis(client):
    return client.get('/api/perfis-financeiros').get_json()['perfis']


def _perfil(client, nome):
    return next(perfil for perfil in _perfis(client) if perfil['nome'] == nome)


def _trocar_perfil(client, nome):
    perfil = _perfil(client, nome)
    response = client.post('/api/perfis-financeiros/ativo', json={'perfil_id': perfil['id']})
    assert response.status_code == 200, response.get_data(as_text=True)
    return perfil


def _documento(
    perfil_id,
    sufixo='doc',
    tipo='NOTA_FISCAL_RECEBIDA',
    categoria='NOTA_FORNECEDOR',
    nome='Fornecedor XPTO',
    ano=2026,
    com_arquivo=True,
    nome_arquivo=None,
):
    comprovante = IrComprovante(
        perfil_financeiro_id=perfil_id,
        ano_calendario=ano,
        data_documento=date(2026, 5, 6),
        prestador_nome=nome,
        prestador_cpf_cnpj='00.000.000/0001-91',
        valor=Decimal('123.45'),
        status='VALIDADO',
        hash_arquivo=f'hash-doc-zip-{perfil_id}-{sufixo}',
        observacoes='Observacao do documento',
    )
    db.session.add(comprovante)
    db.session.flush()
    metadata = DocumentoEmpresarialMetadata(
        comprovante_id=comprovante.id,
        perfil_financeiro_id=perfil_id,
        tipo_documental=tipo,
        categoria_documental=categoria,
        orgao_emissor='Orgao emissor',
        data_emissao=date(2026, 5, 6),
        data_validade=date(2027, 5, 6),
        status_documental='VALIDO',
        observacoes='Observacao empresarial',
    )
    db.session.add(metadata)
    if com_arquivo:
        arquivo = IrComprovanteArquivo(
            comprovante_id=comprovante.id,
            nome_arquivo=nome_arquivo or f'{nome}.pdf',
            mime_type='application/pdf',
            tamanho_bytes=12,
            conteudo=b'%PDF-1.4 ZIP',
            hash_arquivo=comprovante.hash_arquivo,
        )
        db.session.add(arquivo)
    db.session.flush()
    return comprovante


def _zip_response(response):
    return ZipFile(BytesIO(response.data))


def _indice_workbook(zip_file):
    return load_workbook(BytesIO(zip_file.read('indice_documentos.xlsx')))


def test_perfil_empresa_gera_zip_com_indice_e_arquivo(client, app):
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        _documento(empresa['id'], sufixo='nf', nome='Fornecedor XPTO')
        db.session.commit()

    response = client.get('/api/ir/documentos-empresa/exportar-zip?ano=2026')
    zip_file = _zip_response(response)
    nomes = zip_file.namelist()

    assert response.status_code == 200
    assert response.mimetype == 'application/zip'
    assert 'indice_documentos.xlsx' in nomes
    assert any(nome.startswith('notas_fiscais_recebidas/') for nome in nomes)
    assert response.headers['Content-Disposition'].find('Documentos_Empresa_2026_Contador.zip') >= 0


def test_perfil_pessoal_recebe_403_controlado(client):
    response = client.get('/api/ir/documentos-empresa/exportar-zip?ano=2026')

    assert response.status_code == 403
    assert response.get_json()['success'] is False
    assert 'perfil Empresa' in response.get_json()['error']


def test_indice_contem_coluna_nome_no_zip_e_linha_por_documento(client, app):
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        _documento(empresa['id'], sufixo='a', nome='Fornecedor A')
        _documento(empresa['id'], sufixo='b', nome='Fornecedor B')
        db.session.commit()

    response = client.get('/api/ir/documentos-empresa/exportar-zip?ano=2026')
    wb = _indice_workbook(_zip_response(response))
    headers = [cell.value for cell in wb['Indice'][1]]

    assert 'Nome no ZIP' in headers
    assert wb['Indice'].max_row == 3
    assert wb['Indice']['C2'].value.startswith('notas_fiscais_recebidas/')


def test_zip_inclui_todos_os_documentos_filtrados_mesmo_com_paginacao_visual(client, app):
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        for indice in range(7):
            _documento(empresa['id'], sufixo=f'pagina-{indice}', nome=f'Fornecedor pagina {indice}')
        db.session.commit()

    response = client.get('/api/ir/documentos-empresa/exportar-zip?ano=2026')
    zip_file = _zip_response(response)
    wb = _indice_workbook(zip_file)
    arquivos = [nome for nome in zip_file.namelist() if nome.startswith('notas_fiscais_recebidas/')]

    assert response.status_code == 200
    assert len(arquivos) == 7
    assert wb['Indice'].max_row == 8


def test_filtro_por_tipo_documental_define_pasta_e_exclui_outros(client, app):
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        _documento(empresa['id'], sufixo='pag', tipo='PAGAMENTO_REALIZADO', categoria='PIX_PAGO', nome='Pix pago')
        _documento(empresa['id'], sufixo='soc', tipo='DOCUMENTO_SOCIETARIO', categoria='CONTRATO_SOCIAL', nome='Contrato social')
        db.session.commit()

    response = client.get('/api/ir/documentos-empresa/exportar-zip?ano=2026&tipo_documental=DOCUMENTO_SOCIETARIO')
    nomes = _zip_response(response).namelist()

    assert any(nome.startswith('societarios/') for nome in nomes)
    assert not any(nome.startswith('pagamentos_realizados/') for nome in nomes)


def test_documento_de_outro_perfil_nao_entra_no_zip(client, app):
    pessoal = _perfil(client, 'Pessoal')
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        _documento(pessoal['id'], sufixo='pessoal', nome='Documento pessoal oculto')
        _documento(empresa['id'], sufixo='empresa', nome='Documento empresa visivel')
        db.session.commit()

    response = client.get('/api/ir/documentos-empresa/exportar-zip?ano=2026')
    wb = _indice_workbook(_zip_response(response))
    documentos = [row[3].value for row in wb['Indice'].iter_rows(min_row=2)]

    assert 'Documento empresa visivel' in documentos
    assert 'Documento pessoal oculto' not in documentos


def test_nome_de_arquivo_e_sanitizado(client, app):
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        _documento(
            empresa['id'],
            sufixo='sanitizado',
            nome='Fornecedor / XPTO: Nota?',
            nome_arquivo='nota original.pdf',
        )
        db.session.commit()

    response = client.get('/api/ir/documentos-empresa/exportar-zip?ano=2026')
    nomes = _zip_response(response).namelist()
    arquivo_zip = next(nome for nome in nomes if nome.startswith('notas_fiscais_recebidas/'))

    assert '/' in arquivo_zip
    assert not any(caractere in arquivo_zip.split('/')[-1] for caractere in '\\:*?"<>|')
    assert 'Fornecedor_XPTO_Nota' in arquivo_zip


def test_documento_sem_arquivo_fica_no_indice_com_observacao(client, app):
    empresa = _trocar_perfil(client, 'Empresa')
    with app.app_context():
        _documento(empresa['id'], sufixo='sem-arquivo', nome='Documento sem arquivo', com_arquivo=False)
        db.session.commit()

    response = client.get('/api/ir/documentos-empresa/exportar-zip?ano=2026')
    zip_file = _zip_response(response)
    wb = _indice_workbook(zip_file)

    assert response.status_code == 200
    assert zip_file.namelist() == ['indice_documentos.xlsx']
    assert wb['Indice']['C2'].value == 'Arquivo nao incluido'
    assert 'Arquivo nao encontrado ou indisponivel' in wb['Indice']['P2'].value


def test_zip_sem_documentos_retorna_erro_controlado(client):
    _trocar_perfil(client, 'Empresa')

    response = client.get('/api/ir/documentos-empresa/exportar-zip?ano=1999')

    assert response.status_code == 400
    assert response.get_json()['error'] == 'Nenhum documento encontrado para os filtros selecionados.'


def test_frontend_contem_chamada_zip_e_sem_url_externa():
    template = Path('frontend/templates/imposto_renda.html').read_text(encoding='utf-8')
    js = Path('frontend/static/js/imposto_renda.js').read_text(encoding='utf-8')

    assert 'ir-doc-btn-zip' in template
    assert 'Gerar pacote ZIP para contador' in template
    assert '/api/ir/documentos-empresa/exportar-zip' in js
    assert 'http://' not in js[js.find('exportar-zip') - 500:js.find('exportar-zip') + 500]
    assert 'https://' not in js[js.find('exportar-zip') - 500:js.find('exportar-zip') + 500]
